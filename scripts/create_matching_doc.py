"""Create the reviewed human-job matching Word deliverable."""

from __future__ import annotations

import json
import sys
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

PROJECT_ROOT = Path(__file__).resolve().parents[1]
SUBMIT = PROJECT_ROOT / "submit"
DATA = PROJECT_ROOT / "data/processed/jobs_cleaned.parquet"
DATA_JSONL = PROJECT_ROOT / "data/processed/jobs_cleaned.jsonl"
DEMO = PROJECT_ROOT / "data/processed/matching_demo.json"
sys.path.insert(0, str(PROJECT_ROOT / "src"))
sys.path.insert(0, str(PROJECT_ROOT / "scripts"))

from create_submission_docs import (  # noqa: E402
    add_body,
    add_bullets,
    add_page_break,
    add_table,
    add_title,
    configure_document,
    set_run_font,
)


def _code_paragraph(doc: Document, text: str, size: float = 8.8) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    set_run_font(r, name="Consolas", size=size, color="17324D")


def _job_record_count() -> int:
    """Read the Parquet count when available, with a read-only JSONL fallback for doc builds."""
    try:
        return int(len(pd.read_parquet(DATA)))
    except ImportError:
        with DATA_JSONL.open("r", encoding="utf-8") as handle:
            return sum(1 for line in handle if line.strip())


def build_doc() -> Path:
    job_count = _job_record_count()
    demo = json.loads(DEMO.read_text(encoding="utf-8"))
    example_results = demo.get("results", [])

    doc = Document()
    configure_document(doc, "人岗匹配算法设计与实现")
    add_title(
        doc,
        "人岗匹配算法设计与实现",
        "简历输入、文本处理与可解释岗位排序",
        "课程项目：基于全球远程招聘岗位大数据的岗位质量评分与人岗匹配研究  |  审核日期：2026-09-10",
    )
    add_body(
        doc,
        "本报告是第 2 周人岗匹配模块的审核版，基于 Himalayas 公开 Jobs API 的 12,000 条全球远程岗位清洗结果，说明简历如何进入系统、如何统一处理、如何计算匹配分，以及结果如何回溯到岗位来源。当前实现为本地 TF-IDF + 余弦相似度基线，不调用外部大模型。",
    )
    add_body(
        doc,
        "重要口径：match_score 是岗位库内的相对排序参考分，不是企业信誉、岗位质量分或录用概率；当前匹配结果中的 quality_level 仅代表岗位信息的透明度与完整度。",
        bold_lead="重要口径：",
    )

    doc.add_heading("一 目标、范围与数据契约", level=1)
    add_body(
        doc,
        "业务目标是将候选人简历与岗位库进行可解释排序，帮助求职者先查看文本和技能更相关的岗位，再根据技能缺口、经验差异和地区限制做人工判断。模型不替代招聘决策，也不推断候选人的人格、信誉或录用结果。",
    )
    add_table(
        doc,
        ["对象", "实际输入", "核心处理", "结果"],
        [
            ["简历文本", "调用方传入的纯文本（可由页面或 API 收集）", "清洗、jieba 分词、技能与经验识别", "标准文本和候选人特征"],
            ["PDF 简历", "用户主动提供的本地 PDF", "pypdf 逐页提取后走同一流程", "文本提取状态和匹配结果"],
            ["岗位库", "jobs_cleaned.parquet", "标题、标签、类别、描述联合表示", "按 match_score 排序的 Top-k"],
            ["追溯字段", "job_id、source_url 等", "随结果保留，不改写来源", "可回到原始岗位页面核验"],
        ],
        [1.0, 1.65, 2.35, 1.9],
    )
    add_body(
        doc,
        f"数据范围核对：岗位库包含 {job_count:,} 条记录，city 均保持为 Remote；每条记录继续保留 source_url、crawl_time 和 data_origin。匹配模块只读取 data/processed/jobs_cleaned.parquet，不重新联网、不修改 data/raw/ 或 data/raw_responses/。",
    )

    doc.add_heading("二 简历输入处理：从用户输入到可匹配文本", level=1)
    add_body(
        doc,
        "系统支持两条输入路径，但两条路径在得到文本后进入完全相同的清洗、分词和特征识别流程。这样可以避免“粘贴文本”和“PDF 文本”因预处理不一致而产生不可解释的排序差异。",
    )
    add_table(
        doc,
        ["输入路径", "具体处理", "状态与边界"],
        [
            ["简历文本", "调用方传入字符串，直接传入 match_jobs(resume_text, jobs)", "空字符串不执行排序，函数返回空列表；调用方可据此提示补充内容"],
            ["文本文件", "命令行读取 UTF-8 文本，再按粘贴文本路径处理", "文件不存在、权限不足时作为输入错误"],
            ["文本型 PDF", "PdfReader 读取文件；对每一页调用 page.extract_text()，以换行拼接", "保留页间换行，提取后进入同一清洗流程"],
            ["扫描型 PDF", "pypdf 通常只能得到空文本", "函数返回空文本，调用方可提示未检测到可复制文本；当前不伪造 OCR"],
        ],
        [1.25, 3.35, 2.3],
        font_size=8.9,
    )
    add_body(doc, "PDF 提取的最小实现如下，扫描件没有文本层时 `text` 为空，这是当前版本明确的能力边界：")
    _code_paragraph(
        doc,
        'reader = PdfReader(str(path)); text = "\\n".join(page.extract_text() or "" for page in reader.pages).strip()',
        size=8.5,
    )
    add_body(
        doc,
        "输入处理不会把简历原文写回岗位库；岗位库只保存已采集岗位及其分析字段。简历文本在当前演示中以内存变量参与一次匹配，输出仅包含岗位结果、分数和解释。",
    )

    doc.add_heading("三 统一文本预处理与技能归一化", level=1)
    add_body(
        doc,
        "简历与岗位采用同一套标准化流程。岗位侧组合 job_title、skills_normalized、broad_category/job_category 和 description_clean；简历侧对用户文本执行相同清洗，保证二者可以进入同一个 TF-IDF 空间。",
    )
    add_table(
        doc,
        ["步骤", "实现口径", "对匹配的作用"],
        [
            ["1. 清洗", "clean_text：HTML 转纯文本、移除 URL、合并空白、清理大部分标点", "减少网页噪声，保留 Python、C++、.NET 等技能符号"],
            ["2. 分词", "tokenize_text：jieba 处理中英文混合文本，过滤基础停用词", "形成可用于 TF-IDF 的 token 序列"],
            ["3. 技能抽取", "extract_skills：按别名表识别标准技能集合", "计算命中技能、缺失技能和集合重合率"],
            ["4. 经验识别", "extract_years：识别 X years、X yrs、X 年等表达", "与岗位经验代理区间比较"],
            ["5. 类别推断", "infer_category：按类别关键词计数，无法识别则保留未知", "作为类别分项的结构化补充"],
        ],
        [1.0, 3.35, 2.55],
        font_size=8.8,
    )
    add_table(
        doc,
        ["原始写法", "规范写法", "示例用途"],
        [
            ["Golang / Go", "go", "避免同一语言被拆成两个技能"],
            ["Sklearn / scikit-learn", "scikit-learn", "提高机器学习岗位的命中一致性"],
            ["AWS / Amazon Web Services", "aws", "统一云平台表达"],
            ["RESTful / REST API", "rest api", "统一接口开发表达"],
        ],
        [2.1, 1.8, 3.0],
        font_size=8.8,
    )
    add_body(
        doc,
        "技能标签来自平台分类体系，可能混入角色词或职能词。因此系统把标题、平台技能标签、类别和岗位描述联合用于岗位技能识别，不把每个标签机械地当作候选人能力证明。",
    )

    doc.add_heading("四 匹配算法与评分公式", level=1)
    add_body(
        doc,
        "岗位文本和简历文本使用 TfidfVectorizer(ngram_range=(1, 2), sublinear_tf=True) 生成稀疏向量，并以余弦相似度衡量词面相关性。技能分项再融合规范技能集合的重合率，最终得到 0 至 100 的排序分。",
    )
    _code_paragraph(doc, "skill_score = 100 × [0.70 × cosine_similarity + 0.30 × matched_skill_count / job_skill_count]", size=9.0)
    _code_paragraph(doc, "match_score = 0.55 × skill_score + 0.20 × category_score + 0.15 × experience_score + 0.10 × location_score", size=9.0)
    add_table(
        doc,
        ["分项", "权重", "当前规则", "缺失值处理"],
        [
            ["技能相关性", "55%", "TF-IDF 余弦 + 技能集合重合率", "无技能标签时重合率为 0，描述仍参与向量相似度"],
            ["岗位类别", "20%", "类别一致 100，不一致 0", "简历或岗位类别未知时为中性值 50"],
            ["经验/职级", "15%", "落在岗位代理区间为 100，按差距递减", "任一侧缺失时为中性值 50"],
            ["地点限制", "10%", "偏好地区命中 100，不命中 0；Remote 为 100", "岗位地区未知或候选人未给偏好时为 50"],
            ["薪资", "展示项", "保留原币种和周期，不参与总分", "缺失显示“未披露”，不填 0"],
        ],
        [1.3, 0.75, 2.55, 2.3],
        font_size=8.6,
    )
    add_body(
        doc,
        "薪资不进入当前总分是为了避免 6,815 条未披露薪资、不同币种以及 hourly/monthly/annual 周期差异造成机械降权。它仍会作为岗位原始信息展示，供用户自行判断。",
    )

    doc.add_heading("五 输出字段与可解释性", level=1)
    add_body(doc, "每条 Top-k 结果保留分数、依据和来源字段，便于复核而不是只给出一个黑盒排名。")
    add_table(
        doc,
        ["输出字段", "含义", "是否可回溯"],
        [
            ["match_score", "综合排序参考分（0–100）", "可回溯到四个分项"],
            ["matched_skills / missing_skills", "规范技能交集与岗位技能缺口", "可回溯到简历词项和岗位文本"],
            ["cosine_similarity", "简历与岗位联合文本的词面相似度", "可复算，非深度语义分"],
            ["quality_level", "岗位信息透明度与完整度等级", "来自岗位清洗与评分字段"],
            ["source_url", "岗位来源页面", "可打开核验；不代表平台背书"],
            ["crawl_time / data_origin", "采集时间与数据来源标识", "随岗位记录保留"],
        ],
        [1.65, 3.25, 2.0],
        font_size=8.8,
    )
    add_body(
        doc,
        "规则化解释会直接列出综合分、技能分项、命中/缺失技能、类别、经验、地点和薪资参考，不生成“适合录用”“能力优秀”等无法从输入验证的候选人评价。",
    )

    add_page_break(doc)
    doc.add_heading("六 脱敏简历演示与实际 Top-k", level=1)
    resume = "Data analyst with 4 years of experience.\nSkills: Python, SQL, pandas, numpy, Tableau, Excel, statistics.\nExperience: built data pipelines, dashboards, A/B testing reports and machine learning features.\nPreferred location: United States or Canada."
    add_body(doc, "脱敏简历输入：")
    _code_paragraph(doc, resume, size=8.8)
    add_body(doc, "系统识别到 Python、SQL、pandas、numpy、Tableau、Excel、statistics、machine learning 等技能，经验为 4 年，类别线索为 Data Science。以下结果来自当前岗位库中的示例输出：")
    rows = []
    for item in example_results[:5]:
        rows.append(
            [
                str(item.get("job_title", ""))[:36],
                str(item.get("company_name", ""))[:18],
                f"{item.get('match_score', 0):.2f}",
                "、".join(item.get("matched_skills", [])[:4]) or "无",
                "、".join(item.get("missing_skills", [])[:3]) or "无",
            ]
        )
    add_table(doc, ["岗位", "公司", "匹配分", "命中技能", "主要缺失技能"], rows, [2.55, 1.25, 0.7, 1.25, 1.15], font_size=8.0)
    if example_results:
        add_body(doc, "首位结果解释：" + str(example_results[0].get("explanation", "")))
        add_body(doc, "source_url：" + str(example_results[0].get("source_url", "")))
    add_body(doc, "演示解读：分数用于在本岗位库内排序；命中技能和缺失技能用于解释推荐依据；薪资与来源链接仍需用户自行核验。该结果不能外推为录用概率。")

    doc.add_heading("七 实现入口、复现与边界测试", level=1)
    add_table(
        doc,
        ["模块", "入口/函数", "验证内容"],
        [
            ["PDF 输入", "extract_pdf_text(path)", "逐页提取；无文本层时返回空文本边界"],
            ["技能识别", "extract_skills(text)", "别名归一化和技能集合交集"],
            ["匹配排序", "match_jobs(resume_text, jobs, top_k)", "稳定排序、分项分数和 Top-k"],
            ["解释生成", "build_explanation(result)", "命中/缺失技能、分项和薪资参考"],
            ["命令行", "python -m job_analysis.matching resume.txt --top-k 10", "文本文件和 PDF 两类输入"],
        ],
        [1.35, 3.0, 2.55],
        font_size=8.7,
    )
    _code_paragraph(doc, "PYTHONPATH=src .venv/bin/python -m job_analysis.matching resume.pdf --top-k 10 --output data/processed/matching_demo.json", size=8.3)
    add_table(
        doc,
        ["边界场景", "系统行为", "解释"],
        [
            ["简历内容为空", "match_jobs 返回空列表；调用方据此提示补充输入", "避免空向量产生误导排名"],
            ["PDF 没有文本层", "extract_pdf_text 返回空字符串；调用方据此提示", "不把 OCR 冒充为已完成能力"],
            ["无技能词", "技能重合率为 0，其他未知分项保留中性值", "不把抽取失败直接当作能力不足"],
            ["岗位类别/经验缺失", "对应分项为 50", "未知不等于不匹配"],
            ["地区未提供", "地点分项为 50；Remote 岗位按规则为 100", "保持全球远程口径，不改写为厦门"],
            ["薪资缺失/多币种", "显示未披露或原始币种，不进总分", "避免填 0 或假设汇率"],
        ],
        [1.55, 2.95, 2.4],
        font_size=8.6,
    )
    add_body(doc, "已通过 Week 2 模块测试的相关检查包括：技能别名归一化、Top-k 检索/匹配、Remote 地点分项、未知字段中性处理、排序稳定性和来源字段保留。")

    doc.add_heading("八 审核结论与局限", level=1)
    add_bullets(
        doc,
        [
            "输入处理已明确覆盖粘贴文本、文本文件和文本型 PDF；扫描 PDF 的 OCR 尚未实现，文档按边界处理并给出用户提示。",
            "匹配分是可解释排序参考；岗位质量分是信息透明度与完整度，两者不表示企业信誉或录用概率。",
            "TF-IDF 依赖词面共现，对同义表达、上下文和熟练程度的理解有限；后续可增加句向量召回，但应保留当前方法作为可解释基线。",
            "技能标签可能混入角色词，经验区间来自来源职级映射，均属于分析代理；需要人工标注或投递反馈才能评估真实有效性。",
            "当前没有人工相关性金标准、点击率或录用结果，因此不报告虚假的准确率、召回率或录用预测指标。",
        ],
    )
    add_body(doc, "审核结论：人岗匹配模块的输入处理、算法公式、输出语义和异常边界已与当前代码对齐；可作为第 2 周提交材料，后续扩展 OCR、语义向量或人工相关性标注时再更新版本。")

    output = SUBMIT / "人岗匹配文档.docx"
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    return output


if __name__ == "__main__":
    print(build_doc())

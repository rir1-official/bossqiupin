"""Create the two Week 3 submission documents from verified project results."""

from __future__ import annotations

import json
from pathlib import Path
from typing import Iterable, Sequence

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SUBMIT = ROOT / "submit"
REPORTS = ROOT / "reports/week3"
TIMINGS_PATH = REPORTS / "api_test_timings_optimized.json"

FONT_CN = "PingFang SC"
ACCENT = "1D4ED8"
HEADER_FILL = "17324D"
LIGHT_FILL = "EEF4F8"
BORDER = "D9E1E8"


def set_run_font(run, size: float | None = None, bold: bool | None = None, color: str | None = None) -> None:
    run.font.name = FONT_CN
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    run._element.rPr.rFonts.set(qn("w:ascii"), FONT_CN)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_CN)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "5")
        node.set(qn("w:color"), BORDER)


def repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def keep_with_next(paragraph) -> None:
    paragraph.paragraph_format.keep_with_next = True


def configure_document(doc: Document, title: str) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    for style_name, size, bold in (
        ("Normal", 10.5, False),
        ("Title", 24, True),
        ("Heading 1", 16, True),
        ("Heading 2", 12.5, True),
        ("Heading 3", 11, True),
    ):
        style = doc.styles[style_name]
        style.font.name = FONT_CN
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT_CN)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_CN)
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor(0, 0, 0)
    normal = doc.styles["Normal"]
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.space_after = Pt(6)
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(13)
    doc.styles["Heading 1"].paragraph_format.space_after = Pt(7)
    doc.styles["Heading 2"].paragraph_format.space_before = Pt(9)
    doc.styles["Heading 2"].paragraph_format.space_after = Pt(5)

    header = section.header.paragraphs[0]
    header.text = title
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        set_run_font(run, 8.5, color="6B7280")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("第 ")
    set_run_font(run, 8.5, color="6B7280")
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_end])
    run2 = footer.add_run(" 页")
    set_run_font(run2, 8.5, color="6B7280")


def add_cover(doc: Document, title: str, subtitle: str, meta: Sequence[tuple[str, str]]) -> None:
    doc.add_paragraph()
    doc.add_paragraph()
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    set_run_font(r, 24, True, "000000")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(subtitle)
    set_run_font(r, 12, False, ACCENT)

    doc.add_paragraph()
    table = doc.add_table(rows=len(meta), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for idx, (label, value) in enumerate(meta):
        left, right = table.rows[idx].cells
        left.width = Cm(4.0)
        right.width = Cm(9.5)
        left.text = label
        right.text = value
        set_cell_shading(left, LIGHT_FILL)
        for run in left.paragraphs[0].runs:
            set_run_font(run, 10, True, "334155")
        for run in right.paragraphs[0].runs:
            set_run_font(run, 10, False, "111827")
        left.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        left.vertical_alignment = right.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_margins(left, 120, 130, 120, 130)
        set_cell_margins(right, 120, 130, 120, 130)
    set_table_borders(table)
    doc.add_page_break()


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_heading(text, level=level)
    keep_with_next(p)


def add_body(doc: Document, text: str, bold_lead: str | None = None) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0.74)
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, 10.5, True)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2, 10.5)
    else:
        r = p.add_run(text)
        set_run_font(r, 10.5)


def add_bullets(doc: Document, items: Iterable[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Cm(0.7)
        p.paragraph_format.first_line_indent = Cm(-0.35)
        r = p.add_run(item)
        set_run_font(r, 10.5)


def add_table(
    doc: Document,
    headers: Sequence[str],
    rows: Sequence[Sequence[object]],
    widths_cm: Sequence[float] | None = None,
    font_size: float = 9.0,
) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    repeat_header(table.rows[0])
    for col, header in enumerate(headers):
        cell = table.rows[0].cells[col]
        cell.text = str(header)
        set_cell_shading(cell, HEADER_FILL)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_margins(cell, 100, 90, 100, 90)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in cell.paragraphs[0].runs:
            set_run_font(run, font_size, True, "FFFFFF")
    for row_idx, values in enumerate(rows):
        cells = table.add_row().cells
        for col, value in enumerate(values):
            cell = cells[col]
            cell.text = str(value)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, 90, 90, 90, 90)
            if row_idx % 2:
                set_cell_shading(cell, "F8FAFC")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    set_run_font(run, font_size)
        if widths_cm:
            for col, width in enumerate(widths_cm):
                cells[col].width = Cm(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_landscape_section(doc: Document):
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    section.orientation = 1
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.6)
    section.right_margin = Cm(1.6)
    return section


def load_timings() -> dict:
    return json.loads(TIMINGS_PATH.read_text(encoding="utf-8"))


def create_test_report() -> Path:
    timings = load_timings()
    warm = {item["case"]: item for item in timings["warm_cases"]}
    cold = timings["cold_initialization"]["seconds"]

    doc = Document()
    configure_document(doc, "系统测试与优化报告")
    add_cover(
        doc,
        "系统测试与优化报告",
        "基于全球远程招聘岗位大数据的岗位质量评分与人岗匹配研究",
        [
            ("测试阶段", "Week 3 系统集成与优化"),
            ("测试日期", "2026 年 9 月 19 日"),
            ("数据范围", "Himalayas 全球远程岗位 12,000 条"),
            ("实施范围", "FastAPI 接口、Streamlit 前端、简历解析、人岗匹配与 Agent 工作流"),
        ],
    )

    add_heading(doc, "1 报告结论")
    add_body(doc, "本轮测试围绕系统的正常、边界和异常三条路径展开。Week 3 接口回归测试共 13 项，结果为 13 项通过、0 项失败，总耗时 14.083 秒。测试覆盖服务健康检查、文本及 TXT 简历输入、Top-k 边界、聚类查询、岗位评分、本地 Agent 路由以及多类错误输入。")
    add_body(doc, f"人岗匹配已完成索引复用优化。同一测试环境中，优化前的单次匹配中位数约为 13.1534 秒；优化后冷启动首次请求耗时 {cold:.4f} 秒，索引预热后 Top-5 请求中位数为 {warm['warm_match_top_k_5']['median_seconds']:.4f} 秒。冷启动成本仍然存在，但连续演示时的等待已显著降低。")

    add_heading(doc, "2 测试范围与环境")
    add_table(
        doc,
        ["项目", "实际配置"],
        [
            ("操作系统", "macOS 本地环境"),
            ("Python", "项目 .venv312，Python 3.12"),
            ("接口框架", "FastAPI + TestClient"),
            ("前端", "Streamlit 网页端"),
            ("数据", "data/processed/jobs_cleaned.parquet，12,000 条"),
            ("数据约束", "只读本地清洗数据；不重新采集；不修改 raw 目录"),
        ],
        [3.5, 12.5],
        9.5,
    )

    add_heading(doc, "3 接口与功能覆盖")
    add_table(
        doc,
        ["方法", "路径", "主要功能", "关键检查"],
        [
            ("GET", "/api/health", "服务及数据健康检查", "状态为 ok，数据量为 12,000"),
            ("POST", "/api/match", "文本简历人岗匹配", "Top-k、分项分、source_url 与质量分"),
            ("POST", "/api/match/upload", "TXT/PDF 简历上传", "文件类型、空文件、地点偏好"),
            ("POST", "/api/score", "岗位信息质量分查询", "缺少参数与未知岗位"),
            ("POST", "/api/cluster", "K-Means 聚类摘要", "聚类 ID 与业务名称"),
            ("POST", "/api/retrieve", "FAISS+BGE 岗位检索", "Top-k 证据与本地回退"),
            ("POST", "/api/chat", "Agent 工具路由", "本地模式不发起外部模型请求"),
        ],
        [1.3, 3.3, 5.0, 6.2],
        8.6,
    )

    add_landscape_section(doc)
    add_heading(doc, "4 详细测试用例")
    test_rows = [
        ("W3-01", "正常", "GET /health", "无", "200；status=ok；12,000 条", "通过"),
        ("W3-02", "正常", "POST /match", "Python、SQL、pandas，top_k=2", "返回匹配岗位、quality_score、source_url", "通过"),
        ("W3-03", "异常", "POST /match", "空字符串", "Pydantic 校验拒绝，422", "通过"),
        ("W3-04", "异常", "POST /match", "只包含空格、换行、Tab", "业务校验拒绝，400", "通过"),
        ("W3-05", "边界", "POST /match", "top_k=1 和 top_k=20", "返回条数在边界内", "通过"),
        ("W3-06", "异常", "POST /match", "top_k=0 和 top_k=21", "越界参数被拒绝，422", "通过"),
        ("W3-07", "异常", "POST /score", "缺参；未知 job_id", "分别返回 400 和 404", "通过"),
        ("W3-08", "正常", "POST /cluster", "cluster_id=0", "返回聚类 ID 和业务名称", "通过"),
        ("W3-09", "正常", "POST /chat", "use_real_model=false", "model_call=false；本地规则路由", "通过"),
        ("W3-10", "正常", "POST /match/upload", "TXT 简历、top_k=1、地点偏好", "复用匹配流水线并返回 location_score", "通过"),
        ("W3-11", "异常", "POST /match/upload", "空 TXT 文件", "返回 400 和可定位信息", "通过"),
        ("W3-12", "异常", "POST /match/upload", "top_k=21", "返回 400", "通过"),
        ("W3-13", "异常", "POST /match/upload", "DOCX 文件", "非 PDF/TXT 类型被拒绝，400", "通过"),
    ]
    add_table(doc, ["编号", "路径", "接口", "输入与条件", "预期结果", "实际结果"], test_rows, [1.6, 1.8, 4.0, 7.0, 8.5, 2.0], 8.4)
    add_body(doc, "执行命令：./scripts/python.sh -m unittest tests.test_week3_api -v。本轮结果为 Ran 13 tests in 14.083s，OK。")

    section = doc.add_section(WD_SECTION.NEW_PAGE)
    section.orientation = 0
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    add_heading(doc, "5 Bug 修复与性能优化")
    add_table(
        doc,
        ["问题", "原因", "修复方式", "验证"],
        [
            ("匹配请求等待约 13 秒", "每次请求重复拟合 12,000 条岗位的 TF-IDF 矩阵", "LocalJobAgent 初始化时构建 MatchingIndex，后续请求只转换简历向量", "预热后中位数 0.1705–0.2108 秒"),
            ("纯空白简历可进入业务层", "min_length=1 不会过滤空格和换行", "strip 后增加显式空值校验", "新增 W3-04 回归测试并通过"),
            ("性能结果缺少可复跑记录", "旧文件只有优化前测量", "新增 benchmark_week3_api.py 并输出 JSON", "api_test_timings_optimized.json 已生成"),
        ],
        [3.2, 4.3, 5.6, 3.3],
        8.7,
    )

    add_heading(doc, "6 响应时间对比")
    add_table(
        doc,
        ["测量场景", "运行次数", "中位数或单次耗时", "说明"],
        [
            ("优化前普通匹配", "3", "13.1534 秒（中位数）", "每次重建 TF-IDF 索引"),
            ("优化后冷启动首次匹配", "1", f"{cold:.4f} 秒", "包含读取 Parquet、初始化 Agent、构建索引和首次匹配"),
            ("优化后预热 Top-5", "5", f"{warm['warm_match_top_k_5']['median_seconds']:.4f} 秒（中位数）", "复用岗位向量索引"),
            ("优化后预热 Top-1", "5", f"{warm['warm_match_top_k_1']['median_seconds']:.4f} 秒（中位数）", "边界值场景"),
            ("优化后预热 Top-20", "5", f"{warm['warm_match_top_k_20']['median_seconds']:.4f} 秒（中位数）", "最大返回数场景"),
        ],
        [4.5, 2.0, 4.6, 5.3],
        9.0,
    )
    add_body(doc, "测量使用 FastAPI TestClient 和项目 Python 环境。结果反映本机条件，不等同于云服务器 SLA。预热后请求加速主要来自索引复用，冷启动耗时尚未消除。")

    add_heading(doc, "7 Docker 与部署验证")
    add_body(doc, "项目已编写 Dockerfile、docker-compose.yml、requirements-docker.txt 和 .dockerignore。docker compose config --quiet 可正常解析，Compose 包含 api:8000 与 frontend:8501 两个服务，前端依赖 API 健康检查。")
    add_body(doc, "2026 年 9 月 19 日首次构建因 Docker Hub 拉取 python:3.12-slim 超时；随后使用本机已有的 moodtune-pyspark:3.5.3 作为临时构建基础镜像完成了 ARM64 本地镜像构建。docker compose -p bossqiupin up -d 已成功启动 api 和 frontend 两个容器，API healthcheck 为 healthy，/api/health 返回 12,000 条数据，8501 前端返回 HTTP 200。Dockerfile 默认基础镜像仍为 python:3.12-slim。云服务器部署和在线地址尚未完成。")

    add_heading(doc, "8 已知限制与后续工作")
    add_bullets(
        doc,
        [
            "PDF 简历依赖可复制文本层，扫描件尚未接入 OCR。",
            "RAG 首次加载 BGE 模型可能较慢，运行进程会缓存检索器。",
            "scikit-learn 持久化对象由 1.6.1 生成、当前运行时为 1.9.1，测试出现版本警告，后续应统一版本并重建索引。",
            "Starlette/httpx 出现弃用提示，当前未影响用例执行，但需在依赖升级时处理。",
            "match_score 是岗位排序参考，不是录用概率；quality_score 表示岗位信息透明度与完整度，不是企业信誉。",
            "Week 3 已完成本地容器验证、测试报告和 AI 使用说明。由于当前没有云平台账号，云服务器部署与在线地址尚未生成；15–20 分钟最终答辩 PPT 与模拟答辩不纳入本次文档交付。",
        ],
    )

    path = SUBMIT / "测试报告.docx"
    doc.save(path)
    return path


def create_ai_report() -> Path:
    doc = Document()
    configure_document(doc, "AI 使用说明与反思报告")
    add_cover(
        doc,
        "AI 使用说明与反思报告",
        "项目设计、AI 辅助范围与质量控制说明",
        [
            ("项目名称", "基于全球远程招聘岗位大数据的岗位质量评分与人岗匹配研究"),
            ("报告日期", "2026 年 9 月 19 日"),
            ("使用原则", "人工主导方案与验收，AI 辅助部分代码、排错和文档生成"),
        ],
    )

    add_heading(doc, "1 使用说明")
    add_body(doc, "本项目使用 AI 作为编程与文档助手。我负责项目选题、业务目标、数据口径、大部分算法逻辑、前端产品定位与视觉风格，并对最终代码、测试结果和报告内容进行检查。AI 辅助实现部分 Python 代码、接口联调、测试脚本、性能排查与文档排版。")
    add_body(doc, "项目中的数据来源为 Himalayas 公开 Jobs API，city=Remote 是研究口径。AI 没有自主更换数据源，也没有在 Week 3 重新采集 12,000 条数据。所有新增功能均基于已清洗的本地岗位数据。")

    add_heading(doc, "2 人工与 AI 分工边界")
    add_table(
        doc,
        ["环节", "我的主要工作", "AI 辅助内容", "最终确认方式"],
        [
            ("研究选题与范围", "确定远程招聘、岗位质量与人岗匹配为主线", "协助整理需求和任务清单", "根据课程表和项目目标人工确认"),
            ("数据口径", "选择 Himalayas、保持 city=Remote，要求保留溯源字段", "辅助检查字段与文档表述", "检查 source_url、crawl_time、data_origin"),
            ("算法逻辑", "设计评分、匹配、聚类、RAG 与 Agent 的功能目标和主要口径", "协助把方案落地为代码，补充参数校验和解释输出", "实际运行、对比结果、人工审阅"),
            ("前端产品设计", "主导导航、页面动线、黑白视觉、间距和文案调整", "根据反馈修改 Streamlit/CSS 代码并调试交互", "浏览器逐页检查和人工试用"),
            ("接口与测试", "确定要求覆盖正常、边界和异常路径", "协助编写 FastAPI 接口、单元测试和性能基准脚本", "13 项用例实际执行通过"),
            ("文档与汇报", "确定报告范围、取舍和表述重点", "协助组织内容、生成 Word/PPT 和检查版式", "人工核对指标、截图、未完成项和语义边界"),
        ],
        [2.6, 5.0, 5.0, 3.8],
        8.4,
    )

    add_heading(doc, "3 AI 辅助的主要工作")
    add_heading(doc, "3.1 部分代码实现", level=2)
    add_body(doc, "AI 辅助把已确认的分析思路转化为可运行代码，例如 K-Means 的 K 值评估、PCA 可视化、TF-IDF 匹配、FAISS+BGE 检索、FastAPI 端点、Streamlit 交互和 Agent 工具 Schema。这些代码不是直接接受，而是通过本地运行、结果对比和人工修改后进入项目。")

    add_heading(doc, "3.2 调试与性能排查", level=2)
    add_body(doc, "在 Agent 响应过慢和接口超时的排查中，AI 辅助跟踪到人岗匹配每次重复构建 12,000 条岗位的 TF-IDF 矩阵。修改为 MatchingIndex 复用后，预热请求中位数降至约 0.17–0.21 秒。我通过实际演示需求确认该优化有必要，并用回归测试检查功能没有被破坏。")

    add_heading(doc, "3.3 文档生成与格式检查", level=2)
    add_body(doc, "AI 辅助生成 Markdown 与 DOCX，并通过 Word 渲染结果检查中文字体、表格换页、文字溢出和页面层级。报告中的测试数量、性能数据和部署状态都来自实际文件或命令输出，没有用 AI 猜测值替代。")

    add_heading(doc, "4 算法与模型使用边界")
    add_table(
        doc,
        ["功能", "实际方法", "需要避免的误解"],
        [
            ("岗位质量评分", "基于信息完整度与透明度的规则标签，并训练分类模型", "不得说成企业信誉或录用概率"),
            ("人岗匹配", "TF-IDF、余弦相似度与技能/类别/经验/地点加权", "match_score 是排序参考，不是录用概率"),
            ("岗位聚类", "K-Means，使用肘部法和轮廓系数选 K，PCA 降维视觉化", "聚类名称是业务解读，不是原始数据标签"),
            ("RAG", "BGE Embedding + FAISS 向量检索；TF-IDF 仅作回退基线", "不把 TF-IDF 称为深度学习；规则金标准不等同人工标注 Recall"),
            ("Agent", "真实 OpenAI-compatible 模型负责理解和工具选择，本地 Python 工具执行", "本地规则路由必须明确标识，不能伪装成真实模型调用"),
        ],
        [3.0, 7.0, 6.4],
        8.8,
    )

    add_heading(doc, "5 AI 使用流程")
    add_bullets(
        doc,
        [
            "先明确任务目标、数据边界、不可修改的目录和需要交付的文件。",
            "要求 AI 先读取现有文件和代码，避免根据旧会话猜测进度。",
            "对算法或交互方案先由我确认，再让 AI 辅助实现具体代码。",
            "每次修改后执行单元测试、接口测试或浏览器检查，不以代码已生成作为完成标准。",
            "报告和 PPT 只写入已核验的数据，对容器构建、云部署和模型调用状态如实标记。",
        ],
    )

    add_heading(doc, "6 质量、隐私与安全控制")
    add_body(doc, "项目不将 API Key 写入代码库或报告。真实模型调用从已有运行时配置中读取凭据，本地工具只访问清洗后的岗位数据。测试使用的简历为演示文本，不向岗位数据目录写入个人简历。")
    add_body(doc, "我对 AI 输出进行三类检查：一是代码能否在项目环境中实际运行；二是结果是否能由 JSON、CSV、图表或测试输出支持；三是文档是否准确区分已完成、未完成和回退方案。")

    add_heading(doc, "7 反思")
    add_body(doc, "AI 在大量重复代码、接口样板、测试用例和文档排版上节省了时间，也能帮助快速缩小 Bug 范围。但它会出现版本滞后、路径错误、接口契约理解不完整、中文字体丢失以及把未验证推测写成事实的风险。因此，AI 更适合担任实现和检查助手，不能替代我对业务目标、算法合理性和产品设计的判断。")
    add_body(doc, "本项目中最明显的例子是前端迭代和性能优化。前端的导航、比例、配色、文案和交互节奏需要我持续试用和提出具体修改；匹配响应速度则必须通过基准测试才能判断是否真正改善。这说明“生成代码”只是开发的一部分，设计决策、测试证据和最终责任仍然由人承担。")

    add_heading(doc, "8 后续改进")
    add_bullets(
        doc,
        [
            "将依赖版本固定在可复现范围，避免持久化索引跨版本警告。",
            "把冷启动索引持久化或改为后台预热，并增加更稳定的性能回归测试。",
            "对 AI 参与的重要修改保留 Prompt 版本、代码差异、测试输出和人工验收记录。",
            "在最终答辩前逐页演练前端流程，并准备可核验的本地结果以应对外部模型或网络波动。",
        ],
    )

    path = SUBMIT / "AI使用说明与反思报告.docx"
    doc.save(path)
    return path


def write_markdown_sources() -> None:
    timings = load_timings()
    warm = {item["case"]: item for item in timings["warm_cases"]}
    test_md = f"""# 系统测试与优化报告

## 结论

Week 3 接口回归测试共 13 项，13 项通过，0 项失败，耗时 14.083 秒。测试覆盖正常、边界和异常三条路径。

## 性能

- 优化前普通匹配中位数：13.1534 秒。
- 优化后冷启动首次匹配：{timings['cold_initialization']['seconds']:.4f} 秒。
- 预热 Top-5 匹配中位数：{warm['warm_match_top_k_5']['median_seconds']:.4f} 秒。
- 预热 Top-1 匹配中位数：{warm['warm_match_top_k_1']['median_seconds']:.4f} 秒。
- 预热 Top-20 匹配中位数：{warm['warm_match_top_k_20']['median_seconds']:.4f} 秒。

## 已修复

1. 复用 MatchingIndex，避免每次请求重建 12,000 条岗位的 TF-IDF 矩阵。
2. 增加纯空白简历校验和回归测试。
3. 新增可复跑的 Week 3 API 性能测量脚本和 JSON 结果。

## 部署状态

Docker 配置文件已完成，Compose 可解析。首次拉取默认基础镜像时出现网络超时，随后使用本机已有基础镜像完成 ARM64 本地构建并启动 API 与前端容器；云服务器部署和在线地址尚未完成。
"""
    ai_md = """# AI 使用说明与反思报告

## 项目分工

项目选题、业务目标、Himalayas 数据源、city=Remote 口径、大部分算法逻辑，以及前端产品定位、导航、交互、配色和视觉风格由项目负责人设计或确认。

AI 辅助完成部分 Python 代码、FastAPI/Streamlit 联调、测试脚本、性能排查、Docker 配置和文档技术生成。关键代码和结果均经过人工审阅和实际测试。

## 使用原则

- 不用 AI 猜测值替代实际指标。
- 不把岗位质量分说成企业信誉或录用概率。
- 不把 TF-IDF 说成深度学习。
- 不在项目文件中写入 API Key 或未授权的隐私数据。
- 代码和文档在本地测试、渲染和人工检查后交付。

## 反思

AI 能加快样板代码、重复测试和文档排版，但也可能引入版本、路径、接口契约和排版错误。业务决策、算法取舍、前端设计、测试验收和最终责任仍然由人承担。
"""
    (REPORTS / "test_report.md").write_text(test_md, encoding="utf-8")
    (REPORTS / "ai_usage_reflection.md").write_text(ai_md, encoding="utf-8")


def main() -> None:
    SUBMIT.mkdir(exist_ok=True)
    REPORTS.mkdir(parents=True, exist_ok=True)
    write_markdown_sources()
    print(create_test_report())
    print(create_ai_report())


if __name__ == "__main__":
    main()
